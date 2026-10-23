import json
import threading
from pathlib import Path

import openpyxl
import pytest

from csibridge_extractor.extractor import cli
from csibridge_extractor.extractor.jobs import Job, save_job
from csibridge_extractor.extractor.runner import run_job
from csibridge_extractor.extractor.source import FakeSource, SourceError
from csibridge_extractor.extractor.writer import OutputLockedError

FIX = Path(__file__).parent / "fixtures" / "fake_model.json"
EN = "–"


def _job(**kw):
    base = dict(frames="1227", joints="10, 11", cases=["Str-I1", "Str-I2"],
                tables={"frame": {"enabled": True, "modes": ["envelope", "correspondence"]},
                        "joint": {"enabled": True, "modes": ["envelope"]}})
    base.update(kw)
    return Job(**base)


def _fixture_count(table, mode, cases, objs):
    d = json.loads(FIX.read_text(encoding="utf-8"))
    ci = 2 if table == "frame" else 1
    return sum(1 for r in d["tables"][table][mode] if r[ci] in cases and r[0] in objs)


def _run_info(path):
    ws = openpyxl.load_workbook(path)["Run info"]
    return [(r[0], r[1]) for r in ws.iter_rows(values_only=True)]


def test_end_to_end_fake(tmp_path):
    out = tmp_path / "o.xlsx"
    src = FakeSource(FIX)
    res = run_job(_job(), src, str(out))
    assert res.complete and src.closed
    assert res.rows == {
        f"Frame Forces {EN} Envelope": _fixture_count("frame", "envelope", {"Str-I1", "Str-I2"}, {"1227"}),
        f"Frame Forces {EN} Corresp.": _fixture_count("frame", "correspondence", {"Str-I1", "Str-I2"}, {"1227"}),
        f"Joint Reactions {EN} Envelope": _fixture_count("joint", "envelope", {"Str-I1", "Str-I2"}, {"10", "11"}),
    }
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Run info", *res.rows]
    for name, n in res.rows.items():
        assert wb[name].max_row == n + 2
    info = dict(_run_info(out))
    assert info["complete"] == "True" and "1227" in info["Job"]


def test_rerun_reports_missing_targets_and_cases(tmp_path):
    job = _job(frames="1227 9999", joints="10", groups=["Nope"], cases=["Str-I1", "Gone", "Str-III"],
               tables={"frame": {"enabled": True, "modes": ["envelope"]},
                       "joint": {"enabled": False, "modes": ["envelope"]}})
    res = run_job(job, FakeSource(FIX), str(tmp_path / "o.xlsx"))
    text = "\n".join(res.warnings)
    assert "9999" in text and "Nope" in text and "Gone" in text and "Str-III" in text
    assert res.complete
    warns = [v for k, v in _run_info(tmp_path / "o.xlsx") if k == "Warning"]
    assert any("9999" in w for w in warns) and any("Str-III" in w for w in warns)


def test_locked_output_fails_fast(tmp_path):
    calls = []

    class Spy(FakeSource):
        def __getattribute__(self, name):
            if name in ("model_info", "list_cases", "model_lists", "units_row", "iter_rows", "notes"):
                calls.append(name)
            return super().__getattribute__(name)

    src = Spy(FIX)
    # unwritable target (folder missing) is rejected by the writer's check_writable
    with pytest.raises(OutputLockedError):
        run_job(_job(), src, str(tmp_path / "missing" / "o.xlsx"))
    assert calls == [] and src.closed


def test_cancel_writes_incomplete_workbook(tmp_path):
    cancel = threading.Event()

    class Cancelling(FakeSource):
        def iter_rows(self, table, mode, cases, objects):
            for r in super().iter_rows(table, mode, cases, objects):
                yield r
                cancel.set()

    src = Cancelling(FIX)
    res = run_job(_job(), src, str(tmp_path / "o.xlsx"), cancel=cancel)
    assert not res.complete and src.closed
    info = dict(_run_info(tmp_path / "o.xlsx"))
    assert info["complete"] == "False" and info["Status"] == "INCOMPLETE"
    assert 0 < sum(res.rows.values()) < 8


def test_source_error_on_one_object_continues(tmp_path):
    class Flaky(FakeSource):
        def iter_rows(self, table, mode, cases, objects):
            if objects == ["10"]:
                raise SourceError("boom")
            return super().iter_rows(table, mode, cases, objects)

    job = _job(tables={"frame": {"enabled": False, "modes": ["envelope"]},
                       "joint": {"enabled": True, "modes": ["envelope"]}})
    res = run_job(job, Flaky(FIX, batch_objects=1), str(tmp_path / "o.xlsx"))
    assert res.complete
    assert res.rows[f"Joint Reactions {EN} Envelope"] == _fixture_count("joint", "envelope",
                                                                         {"Str-I1", "Str-I2"}, {"11"})
    assert any("[10]" in w and "boom" in w for w in res.warnings)


def test_default_out_path(tmp_path):
    job = _job(output_dir=str(tmp_path), name_pattern="{model}_x_{stamp}.xlsx")
    res = run_job(job, FakeSource(FIX))
    p = Path(res.path)
    assert p.parent == tmp_path and p.name.startswith("fake_bridge_x_") and p.exists()


def test_cli_exit_codes(tmp_path, capsys, monkeypatch):
    jp = tmp_path / "j.json"
    out = str(tmp_path / "o.xlsx")
    save_job(_job(), jp)
    # fake fixture has a case without results only when requested -> clean run is 0
    assert cli.main(["run", str(jp), "--out", out, "--fake", str(FIX)]) == 0
    save_job(_job(frames="1227 9999"), jp)
    assert cli.main(["run", str(jp), "--out", out, "--fake", str(FIX)]) == 2
    save_job(_job(cases=["Str-III"]), jp)
    assert cli.main(["run", str(jp), "--out", out, "--fake", str(FIX)]) == 1
    save_job(_job(), jp)
    assert cli.main(["run", str(jp), "--out", str(tmp_path / "nodir" / "o.xlsx"), "--fake", str(FIX)]) == 1
    from csibridge_extractor.extractor import csi_source

    def no_csi(*a, **k):
        raise csi_source.CsiConnectError(csi_source.connect_message("not_running"))
    monkeypatch.setattr(csi_source, "CsiSource", no_csi)
    capsys.readouterr()
    assert cli.main(["run", str(jp), "--out", out]) == 1
    assert "not running" in capsys.readouterr().err


def _joint_job():
    return _job(tables={"frame": {"enabled": False, "modes": ["envelope"]},
                        "joint": {"enabled": True, "modes": ["envelope"]}})


def test_batches_are_chunked(tmp_path):
    seen = []

    class Rec(FakeSource):
        def iter_rows(self, table, mode, cases, objects):
            seen.append(list(objects))
            return super().iter_rows(table, mode, cases, objects)

    run_job(_job(joints="10 11 12"), Rec(FIX, batch_objects=2), str(tmp_path / "o.xlsx"))
    assert [b for b in seen if b[0] in ("10", "11", "12")] == [["10", "11"], ["12"]]
    seen.clear()
    run_job(_job(joints="10 11 12"), Rec(FIX), str(tmp_path / "o2.xlsx"))
    assert [b for b in seen if b[0] in ("10", "11", "12")] == [["10", "11", "12"]]


def test_output_locked_mid_run_saves_alternate(tmp_path, monkeypatch):
    import os
    from csibridge_extractor.extractor import writer as wmod
    out = tmp_path / "o.xlsx"
    real = os.replace

    def fake_replace(src, dst):
        if str(dst) == str(out):
            raise PermissionError("locked")
        return real(src, dst)

    monkeypatch.setattr(wmod.os, "replace", fake_replace)
    real_check = wmod.check_writable
    n = {"i": 0}

    def late_lock(path):  # unlocked at start, locked by the time Run info is written
        n["i"] += 1
        if n["i"] > 1 and str(path) == str(out):
            raise OutputLockedError(path, "locked")
        return real_check(path)

    monkeypatch.setattr(wmod, "check_writable", late_lock)
    out.write_bytes(b"old")
    res = run_job(_job(), FakeSource(FIX), str(out))
    assert res.path == str(tmp_path / "o_1.xlsx") and Path(res.path).exists()
    assert out.read_bytes() == b"old"
    info = dict(_run_info(res.path))
    assert info["Output file"] == res.path
    assert any("o_1.xlsx" in w for w in res.warnings)


def test_all_replaces_fail_keeps_part_file(tmp_path, monkeypatch):
    from csibridge_extractor.extractor import writer as wmod

    def nope(src, dst):
        raise PermissionError("locked")

    monkeypatch.setattr(wmod.os, "replace", nope)
    with pytest.raises(OutputLockedError):
        run_job(_job(), FakeSource(FIX), str(tmp_path / "o.xlsx"))
    assert [f for f in tmp_path.iterdir() if f.name.startswith(".part_")]


@pytest.mark.parametrize("exc", [ValueError("odd"), RuntimeError("boom")])
def test_aborted_run_writes_full_run_info_and_reraises(tmp_path, exc):
    class Bad(FakeSource):
        def iter_rows(self, table, mode, cases, objects):
            raise exc

    out = tmp_path / "o.xlsx"
    with pytest.raises(type(exc)) as ei:
        run_job(_job(frames="1227 9999"), Bad(FIX), str(out))
    assert ei.value.workbook_path == str(out)
    info = _run_info(out)
    d = dict(info)
    assert d["complete"] == "False" and "Job" in d
    assert any(k == "Warning" and "9999" in v for k, v in info)
    assert any(k == "Warning" and "ABORTED" in v for k, v in info)


def test_writer_error_aborts_with_run_info(tmp_path):
    from csibridge_extractor.extractor.writer import WriterError

    class Huge(FakeSource):
        def iter_rows(self, table, mode, cases, objects):
            yield ("1227", 0.0, "x" * 40000) + (0,) * 10

    out = tmp_path / "o.xlsx"
    with pytest.raises(WriterError):
        run_job(_job(), Huge(FIX), str(out))
    assert dict(_run_info(out))["complete"] == "False"


def test_zero_rows_and_absent_objects_warned(tmp_path):
    # frame 1301 has no fixture rows; with 1227 present it is "absent"
    res = run_job(_job(frames="1227 1301"), FakeSource(FIX), str(tmp_path / "o.xlsx"))
    assert any("absent" in w and "1301" in w for w in res.warnings)
    res = run_job(_job(frames="1301"), FakeSource(FIX, batch_objects=1), str(tmp_path / "p.xlsx"))
    assert any("no rows" in w and "1301" in w for w in res.warnings)


def test_wrong_row_length_warned_not_dropped(tmp_path):
    class Short(FakeSource):
        def iter_rows(self, table, mode, cases, objects):
            return iter([("10", "Str-I1", "c", "Max", 1.0)])

    res = run_job(_joint_job(), Short(FIX), str(tmp_path / "o.xlsx"))
    assert res.rows[f"Joint Reactions {EN} Envelope"] == 1
    assert any("wrong number of cells" in w for w in res.warnings)


def test_notes_are_not_warnings(tmp_path):
    class Noted(FakeSource):
        def notes(self):
            return ["correspondence not provided"]

    res = run_job(_joint_job(), Noted(FIX), str(tmp_path / "o.xlsx"))
    assert res.warnings == [] and res.notes == ["correspondence not provided"]
    assert ("Note", "correspondence not provided") in _run_info(tmp_path / "o.xlsx")


def test_final_progress_event(tmp_path):
    ev = []
    run_job(_joint_job(), FakeSource(FIX), str(tmp_path / "o.xlsx"), progress=ev.append)
    assert ev and ev[-1].get("final")


def test_enabled_must_be_bool_and_bad_pattern(tmp_path):
    from csibridge_extractor.extractor.jobs import load_job
    p = tmp_path / "j.json"
    p.write_text(json.dumps({"tables": {"frame": {"enabled": "no"}}}), encoding="utf-8")
    with pytest.raises(ValueError, match="enabled"):
        load_job(p)
    p.write_text(json.dumps({"name_pattern": "{oops}.xlsx"}), encoding="utf-8")
    with pytest.raises(ValueError, match="name_pattern"):
        load_job(p)


def test_forward_slash_and_relative_out_no_false_move(tmp_path, monkeypatch):
    jp = tmp_path / "j.json"
    save_job(_job(), jp)
    assert cli.main(["run", str(jp), "--out", (tmp_path / "b.xlsx").as_posix(), "--fake", str(FIX)]) == 0
    monkeypatch.chdir(tmp_path)
    assert cli.main(["run", str(jp), "--out", "./d.xlsx", "--fake", str(FIX)]) == 0
    res = run_job(_job(), FakeSource(FIX), "./e.xlsx")
    assert not res.warnings and Path(res.path).is_absolute()


def test_all_replaces_fail_sets_part_path(tmp_path, monkeypatch):
    from csibridge_extractor.extractor import writer as wmod
    monkeypatch.setattr(wmod.os, "replace", lambda s, d: (_ for _ in ()).throw(PermissionError("x")))
    with pytest.raises(OutputLockedError) as ei:
        run_job(_job(), FakeSource(FIX), str(tmp_path / "o.xlsx"))
    assert Path(ei.value.workbook_path).name.startswith(".part_")


def test_close_fallback_is_reported(tmp_path, monkeypatch):
    import os
    from csibridge_extractor.extractor import writer as wmod
    out = tmp_path / "o.xlsx"
    real = os.replace
    monkeypatch.setattr(wmod.os, "replace",
                        lambda s, d: (_ for _ in ()).throw(PermissionError("x")) if str(d) == str(out) else real(s, d))
    res = run_job(_job(), FakeSource(FIX), str(out))
    assert res.path.endswith("o_1.xlsx") and any("o_1.xlsx" in w for w in res.warnings)


def test_csv_output_folder_power_query_layout(tmp_path):
    import csv as _csv
    res = run_job(_job(output_format="csv"), FakeSource(FIX), str(tmp_path / "o.xlsx"))
    folder = Path(res.path)
    assert res.complete and folder.is_dir() and folder.name == "o"
    names = sorted(p.name for p in folder.iterdir())
    assert names == ["Frame Forces - Corresp.csv", "Frame Forces - Envelope.csv",
                     "Joint Reactions - Envelope.csv", "Run info.csv"]
    with open(folder / "Frame Forces - Envelope.csv", encoding="utf-8-sig", newline="") as fh:
        rows = list(_csv.reader(fh))
    assert rows[0][:2] == ["Frame", "Station [ft]"] and rows[0][5] == "P [Kip]"
    assert len(rows) - 1 == _fixture_count("frame", "envelope", ["Str-I1", "Str-I2"], {"1227"})
    float(rows[1][5])  # numbers, not a units row
    info = (folder / "Run info.csv").read_text(encoding="utf-8-sig")
    assert "COMPLETE" in info


def test_csv_does_not_overwrite_existing_folder(tmp_path):
    (tmp_path / "o").mkdir()
    (tmp_path / "o" / "keep.txt").write_text("x")
    res = run_job(_job(output_format="csv"), FakeSource(FIX), str(tmp_path / "o"))
    assert Path(res.path).name == "o_1" and (tmp_path / "o" / "keep.txt").exists()
    assert any("already had files" in w for w in res.warnings)


def test_csv_overwrite_replaces_own_files_only(tmp_path):
    folder = tmp_path / "fixed"
    job = _job(output_format="csv", overwrite=True, name_pattern="{model}.xlsx")
    run_job(job, FakeSource(FIX), str(folder))
    (folder / "my notes.txt").write_text("keep")
    # second run without joints: the stale joint CSV must disappear, user file stays
    job2 = _job(output_format="csv", overwrite=True, name_pattern="{model}.xlsx",
                tables={"frame": {"enabled": True, "modes": ["envelope"]},
                        "joint": {"enabled": False, "modes": ["envelope"]}})
    res = run_job(job2, FakeSource(FIX), str(folder))
    assert Path(res.path) == folder
    assert sorted(p.name for p in folder.iterdir()) == ["Frame Forces - Envelope.csv", "Run info.csv",
                                                        "my notes.txt"]


def test_csv_overwrite_locked_file_fails_before_deleting(tmp_path):
    import ctypes
    from ctypes import wintypes
    folder = tmp_path / "fixed"
    job = _job(output_format="csv", overwrite=True, name_pattern="{model}.xlsx")
    run_job(job, FakeSource(FIX), str(folder))
    target = folder / "Joint Reactions - Envelope.csv"
    k32 = ctypes.WinDLL("kernel32", use_last_error=True)
    k32.CreateFileW.restype = wintypes.HANDLE
    k32.CreateFileW.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID,
                                wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
    h = k32.CreateFileW(str(target), 0x80000000 | 0x40000000, 0, None, 3, 0x80, None)  # share 0, like Excel
    assert h not in (None, ctypes.c_void_p(-1).value)
    try:
        with pytest.raises(OutputLockedError):
            run_job(job, FakeSource(FIX), str(folder))
        assert (folder / "Frame Forces - Envelope.csv").exists()  # nothing deleted
        assert (folder / "Run info.csv").exists()
    finally:
        k32.CloseHandle(h)


def test_overwrite_requires_fixed_name():
    from csibridge_extractor.extractor.jobs import validate
    assert any("{stamp}" in e for e in validate(_job(overwrite=True)))
    assert not validate(_job(overwrite=True, name_pattern="{model}_results.xlsx"))
