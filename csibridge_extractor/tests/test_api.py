import json
import threading
import time
from pathlib import Path

import openpyxl

from csibridge_extractor.api import Api
from csibridge_extractor.extractor.source import FakeSource

FIX = Path(__file__).parent / "fixtures" / "fake_model.json"


def _job(tmp_path, **kw):
    d = dict(frames="1227", joints="10, 11", groups=[], cases=["Str-I1", "Str-I2"],
             tables={"frame": {"enabled": True, "modes": ["envelope", "correspondence"]},
                     "joint": {"enabled": True, "modes": ["envelope"]}},
             output_dir=str(tmp_path), name_pattern="{model}_out.xlsx", open_when_done=False)
    d.update(kw)
    return json.dumps(d)


def _wait_done(api, timeout=20):
    log, end = [], time.time() + timeout
    while time.time() < end:
        p = json.loads(api.poll())
        log += p["log"]
        if p["done"]:
            return p, log
        time.sleep(0.02)
    raise AssertionError("run did not finish")


def test_connect_fake():
    api = Api(fake=str(FIX))
    r = json.loads(api.connect())
    assert "error" not in r
    assert r["model_info"]["version"] == "27.1.0"
    names = {c["name"]: c for c in r["cases"]}
    assert names["Str-I1"]["has_results"] and not names["Str-III"]["has_results"]
    assert set(r["groups"]) == {"GirderA", "Piers"}
    assert r["n_frames"] == 3 and "notes" in r


def test_resolve_counts_and_missing(tmp_path):
    api = Api(fake=str(FIX))
    assert "error" in json.loads(api.resolve(_job(tmp_path)))  # not connected yet
    api.connect()
    r = json.loads(api.resolve(_job(tmp_path, frames="1227, 9999", groups=["GirderA", "Nope"])))
    assert r["counts"]["frames"] >= 1
    assert "9999" in r["missing_frames"]
    assert "Nope" in r["missing_groups"]


def test_run_poll_until_done(tmp_path):
    api = Api(fake=str(FIX))
    api.connect()
    assert json.loads(api.run(_job(tmp_path)))["started"]
    p, log = _wait_done(api)
    assert p["error"] is None and p["done"]
    res = p["result"]
    assert res["complete"] is True and res["rows"] and log
    wb = openpyxl.load_workbook(res["path"], read_only=True)
    assert "Run info" in wb.sheetnames
    assert Path(res["path"]).parent == tmp_path


class _Gated(FakeSource):
    def __init__(self, path, started, gate):
        super().__init__(path, batch_objects=1)
        self.started, self.gate = started, gate

    def iter_rows(self, table, mode, cases, objects):
        self.started.set()
        self.gate.wait(10)
        yield from super().iter_rows(table, mode, cases, objects)


def test_cancel_via_api(tmp_path):
    started, gate = threading.Event(), threading.Event()
    api = Api(source_factory=lambda: _Gated(FIX, started, gate))
    api.connect()
    api.run(_job(tmp_path, frames="1227, 1301"))
    assert started.wait(10)
    json.loads(api.cancel())
    gate.set()
    p, _ = _wait_done(api)
    assert p["result"]["complete"] is False
    assert any("cancel" in w.lower() for w in p["result"]["warnings"])


def test_errors_are_json(tmp_path):
    api = Api(fake=str(FIX))
    for raw in (api.run("not json"), api.resolve("{"), api.poll() and api.run("[]")):
        assert "error" in json.loads(raw)
    assert "error" in json.loads(Api(fake=str(tmp_path / "missing.json")).connect())
    api.connect()
    r = json.loads(api.run(_job(tmp_path, cases=[])))  # validation error
    assert "error" in r
    r = json.loads(api.open_file(str(tmp_path / "nope.xlsx")))
    assert "error" in r


def test_run_failure_reports_error(tmp_path):
    api = Api(fake=str(FIX))
    api.connect()
    api.run(_job(tmp_path, output_dir=str(tmp_path / "no" / "such" / "dir")))
    p, _ = _wait_done(api)
    assert p["error"] and p["result"] is None


def test_save_and_load_job(tmp_path):
    api = Api(fake=str(FIX))
    path = tmp_path / "j.json"
    assert json.loads(api.save_job(_job(tmp_path), str(path)))["saved"]
    r = json.loads(api.load_job(str(path)))
    assert r["job"]["frames"] == "1227" and r["job"]["cases"] == ["Str-I1", "Str-I2"]
    assert "error" in json.loads(api.load_job(str(tmp_path / "x.json")))


def test_preview_output(tmp_path):
    api = Api(fake=str(FIX))
    api.connect()
    r = json.loads(api.preview_output(_job(tmp_path)))
    assert r["path"].endswith("fake_bridge_out.xlsx")


def test_app_loads_html_without_pywebview():
    import sys
    from csibridge_extractor import app
    html = app.get_html()
    assert "<title>CSiBridge Results Extractor</title>" in html and "http" not in html.split("<script>")[1]
    assert "webview" not in sys.modules
