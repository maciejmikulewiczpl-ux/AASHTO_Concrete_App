import math
import sys

import openpyxl
import pytest

from csibridge_extractor.extractor.writer import OutputLockedError, WorkbookWriter, check_writable

HDR = ["Frame", "Station", "Case"]
UNITS = ["Text", "ft", "Text"]
TAB = "Frame Forces – Envelope"


def test_two_header_rows_and_types(tmp_path):
    p = tmp_path / "a.xlsx"
    with WorkbookWriter(p) as w:
        t = w.open_table(TAB, HDR, UNITS)
        t.write(("1227", 8.4167, "=SUM(A1)"))
        t.write(("1228", 0, None))
        assert t.rows_written == 2 and t.tabs == [TAB]
    ws = openpyxl.load_workbook(p)[TAB]
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    assert rows[0] == HDR and rows[1] == UNITS
    assert rows[2] == ["1227", 8.4167, "=SUM(A1)"]
    assert ws.cell(3, 3).data_type == "s"  # text stays text, never a formula
    assert ws.cell(3, 2).data_type == "n"
    assert rows[3][1] == 0 and rows[3][2] is None


def test_continuation_tabs_preserve_every_row(tmp_path):
    p = tmp_path / "b.xlsx"
    with WorkbookWriter(p, row_limit=10) as w:
        t = w.open_table("X", HDR, UNITS)
        for i in range(25):
            t.write((str(i), float(i), "c"))
        assert t.tabs == ["X", "X (2)", "X (3)"] and t.rows_written == 25
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames == ["Run info", "X", "X (2)", "X (3)"]
    seen = []
    for name in wb.sheetnames[1:]:
        rows = [[c.value for c in r] for r in wb[name].iter_rows()]
        assert rows[0] == HDR and rows[1] == UNITS
        seen += [r[0] for r in rows[2:]]
    assert seen == [str(i) for i in range(25)]


def test_run_info_first_tab(tmp_path):
    p = tmp_path / "c.xlsx"
    with WorkbookWriter(p) as w:
        t = w.open_table(TAB, HDR, UNITS)
        t.write(("1", 1.0, "c"))
        w.write_run_info([("Complete", "yes"), ("Rows", "1")])  # written after the data tab
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames == ["Run info", TAB]
    vals = [[c.value for c in r] for r in wb["Run info"].iter_rows()]
    assert vals == [["Complete", "yes"], ["Rows", "1"]]


def test_run_info_tab_exists_without_items_and_nan_safe(tmp_path):
    p = tmp_path / "d.xlsx"
    with WorkbookWriter(p) as w:
        t = w.open_table("T", ["a"], ["u"])
        t.write((math.nan,))
        t.write((5,))
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames[0] == "Run info"
    assert [c.value for c in wb["T"]["A"]] == ["a", "u", "nan", 5]


def test_check_writable_ok_and_missing_folder(tmp_path):
    check_writable(tmp_path / "new.xlsx")
    existing = tmp_path / "e.xlsx"
    existing.write_bytes(b"x")
    check_writable(existing)
    assert existing.read_bytes() == b"x"
    with pytest.raises(OutputLockedError):
        check_writable(tmp_path / "nope" / "f.xlsx")


@pytest.mark.skipif(sys.platform != "win32", reason="exclusive file handle simulation is Windows-only")
def test_check_writable_locked(tmp_path):
    import ctypes
    from ctypes import wintypes

    p = tmp_path / "locked.xlsx"
    p.write_bytes(b"x")
    k32 = ctypes.WinDLL("kernel32", use_last_error=True)
    k32.CreateFileW.restype = wintypes.HANDLE
    k32.CreateFileW.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID,
                                wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
    # GENERIC_READ|GENERIC_WRITE, share mode 0 (like Excel), OPEN_EXISTING
    h = k32.CreateFileW(str(p), 0x80000000 | 0x40000000, 0, None, 3, 0x80, None)
    assert h not in (None, ctypes.c_void_p(-1).value)
    try:
        with pytest.raises(OutputLockedError) as ei:
            check_writable(p)
        assert str(p) in str(ei.value)
    finally:
        k32.CloseHandle(h)
    check_writable(p)


def test_tab_names_within_31_chars(tmp_path):
    p = tmp_path / "f.xlsx"
    long = "A very long table tab name exceeding limit"
    with WorkbookWriter(p, row_limit=2) as w:
        t = w.open_table(long, ["a"], ["u"])
        for i in range(5):
            t.write((i,))
        assert all(len(n) <= 31 for n in t.tabs) and len(set(t.tabs)) == 3
    wb = openpyxl.load_workbook(p)
    assert all(len(n) <= 31 for n in wb.sheetnames) and len(wb.sheetnames) == 4


def test_long_text_data_cell_raises_run_info_splits(tmp_path):
    from csibridge_extractor.extractor.writer import WriterError
    p = tmp_path / "g.xlsx"
    big = "x" * 40000
    with WorkbookWriter(p) as w:
        t = w.open_table("T", ["a"], ["u"])
        with pytest.raises(WriterError, match="T"):
            t.write((big,))
        w.write_run_info([("Note", big), ("End", "ok")])
    vals = [[c.value for c in r] for r in openpyxl.load_workbook(p)["Run info"].iter_rows()]
    assert "".join(v[1] for v in vals[:2]) == big
    assert vals[1][0] == "Note (continued)" and vals[2] == ["End", "ok"]


def test_exception_writes_incomplete_run_info_and_propagates(tmp_path):
    p = tmp_path / "h.xlsx"
    with pytest.raises(KeyError):
        with WorkbookWriter(p) as w:
            w.open_table("T", ["a"], ["u"]).write((1,))
            raise KeyError("boom")
    wb = openpyxl.load_workbook(p)
    text = " ".join(str(c.value) for r in wb["Run info"].iter_rows() for c in r)
    assert "INCOMPLETE" in text and "KeyError" in text


def test_bool_numbers_and_counters(tmp_path):
    p = tmp_path / "i.xlsx"
    with WorkbookWriter(p) as w:
        t = w.open_table("T", ["a", "b"], ["u", "u"])
        t.write((1, True))
        t.write((math.inf, False))
        t.write(("n/a", None))
        assert t.non_finite_cells == 1 and t.non_numeric_cells == 1
    ws = openpyxl.load_workbook(p)["T"]
    assert ws.cell(3, 2).value is True and ws.cell(3, 2).data_type == "b"
    assert ws.cell(3, 1).value == 1


def test_atomic_replace_keeps_old_file_on_failure(tmp_path):
    p = tmp_path / "j.xlsx"
    with WorkbookWriter(p) as w:
        w.open_table("T", ["a"], ["u"]).write((1,))
    good = p.read_bytes()
    w2 = WorkbookWriter(p)
    assert p.read_bytes() == good  # untouched until close
    w2.close()
    assert p.read_bytes() != good
    assert not [f for f in tmp_path.iterdir() if f.name.startswith(".part_")]
